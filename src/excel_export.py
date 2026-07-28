"""Xuất kết quả ra Excel (mục 14 CLAUDE.md).

File có ít nhất 3 sheet: Kết quả, Hóa đơn đầu vào, Thông tin xử lý.
Nếu có lỗi kiểm tra, thêm sheet Lỗi hóa đơn / Lỗi dữ liệu nguồn.
"""

from __future__ import annotations

import datetime as _dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from src.models import AllocationResultLine, DataWarning, InvoiceInputRow, RowError

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL_COLOR = "305496"
BODY_FONT = Font(name="Arial")

RESULT_HEADERS = [
    "STT hóa đơn",
    "Dòng phân bổ",
    "MÃ SP",
    "TÊN SP",
    "SL yêu cầu",
    "SL phân bổ",
    "Giá bán",
    "Mã VAT",
    "TÊN VAT",
    "TÊN KẾ TOÁN",
    "ĐVT",
    "NGUỒN",
    "TỒN CÒN NHÁP TRƯỚC PHÂN BỔ",
    "TỒN CÒN MÔ PHỎNG SAU PHÂN BỔ",
    "%VAT",
    "GIÁ NHÁP",
    "Điểm khớp tên",
    "Lý do chọn",
]


def _style_header_row(ws, row_idx: int, ncols: int) -> None:
    from openpyxl.styles import PatternFill

    fill = PatternFill(start_color=HEADER_FILL_COLOR, end_color=HEADER_FILL_COLOR, fill_type="solid")
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.font = HEADER_FONT
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row_idx + 1, column=1).coordinate
    ws.auto_filter.ref = f"A{row_idx}:{get_column_letter(ncols)}{row_idx}"


def _autosize(ws, ncols: int, max_width: int = 45) -> None:
    for c in range(1, ncols + 1):
        letter = get_column_letter(c)
        max_len = 10
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_width, max_len + 2)


def _write_result_sheet(ws, lines: list[AllocationResultLine], decimal_places: int) -> None:
    ws.append(RESULT_HEADERS)
    _style_header_row(ws, 1, len(RESULT_HEADERS))

    for line in lines:
        ws.append(
            [
                line.stt_hoa_don,
                line.dong_phan_bo,
                line.ma_sp,
                line.ten_sp,
                line.sl_yeu_cau,
                line.sl_phan_bo,
                line.gia_ban,
                line.ma_vat,
                line.ten_vat,
                line.ten_ke_toan,
                line.dvt,
                line.nguon,
                line.ton_con_nhap_truoc,
                line.ton_con_mo_phong_sau,
                line.phan_tram_vat,
                line.gia_nhap,
                line.diem_khop_ten,
                line.ly_do_chon,
            ]
        )

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = BODY_FONT
        row[4].number_format = "#,##0"
        row[5].number_format = "#,##0"
        row[6].number_format = "#,##0"
        row[12].number_format = "#,##0"
        row[13].number_format = "#,##0"
        row[14].number_format = "0.0%"
        row[15].number_format = "#,##0." + "0" * decimal_places
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
        row[9].alignment = Alignment(wrap_text=True, vertical="top")

    _autosize(ws, len(RESULT_HEADERS))


def _write_input_sheet(ws, rows: list[InvoiceInputRow]) -> None:
    headers = ["Mã SP", "Tên SP", "SL", "Giá bán"]
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))
    for r in sorted(rows, key=lambda x: x.order_index):
        ws.append([r.ma_sp, r.ten_sp, r.sl, r.gia_ban])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = BODY_FONT
    _autosize(ws, len(headers))


def _write_info_sheet(
    ws,
    source_file_name: str,
    source_file_hash: str,
    invoice_row_count: int,
    total_sl: float,
    config_summary: dict,
    warnings: list[str],
) -> None:
    ws.append(["Trường", "Giá trị"])
    _style_header_row(ws, 1, 2)
    rows = [
        ("Tên file dữ liệu nguồn", source_file_name),
        ("Hash file (SHA-256)", source_file_hash),
        ("Thời gian xử lý", _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Số dòng hóa đơn", invoice_row_count),
        ("Tổng số lượng yêu cầu", total_sl),
    ]
    for k, v in config_summary.items():
        rows.append((f"Cấu hình: {k}", str(v)))
    for w in warnings:
        rows.append(("Cảnh báo", w))
    for k, v in rows:
        ws.append([k, v])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = BODY_FONT
    _autosize(ws, 2)


def _write_error_sheet(ws, errors: list[RowError]) -> None:
    headers = ["Dòng hóa đơn", "Mã SP", "Tên SP", "Loại lỗi", "Nội dung chi tiết", "Hướng xử lý đề xuất"]
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))
    for e in sorted(errors, key=lambda x: x.order_index):
        ws.append([e.order_index, e.ma_sp, e.ten_sp, e.error_type, e.detail, e.suggestion])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(ws, len(headers))


def _write_data_warning_sheet(ws, warnings: list[DataWarning]) -> None:
    headers = ["Dòng trong file nguồn", "Loại cảnh báo", "Nội dung"]
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))
    for w in warnings:
        ws.append([w.row_index, w.warning_type, w.detail])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = BODY_FONT
    _autosize(ws, len(headers))


def export_result_workbook(
    output_path: str | Path,
    result_lines: list[AllocationResultLine],
    input_rows: list[InvoiceInputRow],
    source_file_name: str,
    source_file_hash: str,
    config_summary: dict,
    decimal_places: int = 2,
    processing_warnings: list[str] | None = None,
    errors: list[RowError] | None = None,
    data_warnings: list[DataWarning] | None = None,
) -> Path:
    wb = Workbook()

    ws_result = wb.active
    ws_result.title = "Kết quả"
    _write_result_sheet(ws_result, result_lines, decimal_places)
    ws_result.page_setup.orientation = "landscape"

    ws_input = wb.create_sheet("Hóa đơn đầu vào")
    _write_input_sheet(ws_input, input_rows)

    total_sl = sum(r.sl for r in input_rows if isinstance(r.sl, (int, float)))
    ws_info = wb.create_sheet("Thông tin xử lý")
    _write_info_sheet(
        ws_info,
        source_file_name,
        source_file_hash,
        len(input_rows),
        total_sl,
        config_summary,
        processing_warnings or [],
    )

    if errors:
        ws_err = wb.create_sheet("Lỗi hóa đơn")
        _write_error_sheet(ws_err, errors)

    if data_warnings:
        ws_dw = wb.create_sheet("Lỗi dữ liệu nguồn")
        _write_data_warning_sheet(ws_dw, data_warnings)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def result_lines_to_tsv(lines: list[AllocationResultLine]) -> str:
    """Sao chép bảng kết quả dưới dạng TSV (không HTML) để dán vào Excel."""

    out_rows = ["\t".join(RESULT_HEADERS)]
    for line in lines:
        values = [
            line.stt_hoa_don,
            line.dong_phan_bo,
            line.ma_sp,
            line.ten_sp,
            line.sl_yeu_cau,
            line.sl_phan_bo,
            line.gia_ban,
            line.ma_vat,
            line.ten_vat,
            line.ten_ke_toan,
            line.dvt,
            line.nguon,
            line.ton_con_nhap_truoc,
            line.ton_con_mo_phong_sau,
            f"{line.phan_tram_vat * 100:.1f}%" if line.phan_tram_vat is not None else "",
            line.gia_nhap,
            line.diem_khop_ten,
            line.ly_do_chon,
        ]
        out_rows.append("\t".join("" if v is None else str(v) for v in values))
    return "\n".join(out_rows)
