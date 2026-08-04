import io

import openpyxl

from src.invoice_file_reader import parse_invoice_file_bytes, parse_invoice_files


def _build_invoice_workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ChiTietHoaDon_HD000001"
    ws.append(
        ["Mã hàng", "Tên hàng", "Đơn vị tính", "Số lượng", "Đơn giá", "Giảm giá %", "Giảm giá", "Giá bán", "Thành tiền"]
    )
    ws.append(["VTHS", "Vòi hoa sen nano SN18", None, 200, 25500.0, 0, None, 25500.0, 5100000.0])
    ws.append(["2026DO45", "Chổi gạt nước LAXA 90cm", None, 100, 13800.0, 0, None, 13800.0, 1380000.0])
    ws.append([None, None, None, None, None, None, "Tổng số lượng", None, 300])
    ws.append([None, None, None, None, None, None, "Tổng cộng", None, 6480000.0])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parses_product_rows_and_skips_footer():
    result = parse_invoice_file_bytes(_build_invoice_workbook_bytes(), "hoa_don.xlsx")

    assert result.errors == []
    assert result.has_invoice_meta is False
    assert result.rows == [
        ["VTHS", "Vòi hoa sen nano SN18", "200", "25500", "", "", "5100000"],
        ["2026DO45", "Chổi gạt nước LAXA 90cm", "100", "13800", "", "", "1380000"],
    ]


def _build_invoice_workbook_with_meta_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ChiTietHoaDon_HD000001"
    ws.append(
        [
            "Mã hóa đơn",
            "Mã hàng",
            "Tên hàng",
            "Số lượng",
            "Giá bán",
            "Ghi chú hàng hóa",
            "Thành tiền",
        ]
    )
    ws.append(["HD000001", "VTHS", "Vòi hoa sen nano SN18", 200, 25500.0, "ghi chú", 5100000.0])
    ws.append(["HD000001", "2026DO45", "Chổi gạt nước LAXA 90cm", 100, 13800.0, "", 1380000.0])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parses_invoice_code_and_amount_when_present():
    result = parse_invoice_file_bytes(_build_invoice_workbook_with_meta_bytes(), "hoa_don.xlsx")

    assert result.errors == []
    assert result.has_invoice_meta is True
    assert result.rows == [
        ["VTHS", "Vòi hoa sen nano SN18", "200", "25500", "ghi chú", "HD000001", "5100000"],
        ["2026DO45", "Chổi gạt nước LAXA 90cm", "100", "13800", "", "HD000001", "1380000"],
    ]


def test_rejects_non_xlsx_file_name():
    result = parse_invoice_file_bytes(b"not really an xlsx", "hoa_don.csv")

    assert result.rows == []
    assert any(".xlsx" in e for e in result.errors)


def test_missing_header_columns_reports_error():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Cột A", "Cột B"])
    ws.append(["x", "y"])
    buffer = io.BytesIO()
    wb.save(buffer)

    result = parse_invoice_file_bytes(buffer.getvalue(), "khong_dung_cot.xlsx")

    assert result.rows == []
    assert any("không tìm thấy cột" in e for e in result.errors)


def test_parse_invoice_files_merges_multiple_files():
    file_bytes = _build_invoice_workbook_bytes()

    rows, errors, has_invoice_meta = parse_invoice_files(
        [(file_bytes, "hd1.xlsx"), (file_bytes, "hd2.xlsx")]
    )

    assert len(rows) == 4
    assert errors == []
    assert has_invoice_meta is False


def test_parse_invoice_files_flags_invoice_meta_only_when_all_files_have_it():
    with_meta = _build_invoice_workbook_with_meta_bytes()
    without_meta = _build_invoice_workbook_bytes()

    rows, errors, has_invoice_meta = parse_invoice_files(
        [(with_meta, "hd1.xlsx"), (with_meta, "hd2.xlsx")]
    )
    assert len(rows) == 4
    assert errors == []
    assert has_invoice_meta is True

    rows_mixed, errors_mixed, has_invoice_meta_mixed = parse_invoice_files(
        [(with_meta, "hd1.xlsx"), (without_meta, "hd2.xlsx")]
    )
    assert len(rows_mixed) == 4
    assert errors_mixed == []
    assert has_invoice_meta_mixed is False
