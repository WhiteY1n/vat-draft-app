from pathlib import Path

import openpyxl

from src.excel_reader import read_and_validate_workbook

SAMPLE_PATH = Path(__file__).resolve().parent.parent / "data" / "sample" / "Toi_uu_nhap_VAT_mau.xlsx"


def test_reads_real_sample_file_successfully():
    assert SAMPLE_PATH.exists(), "Thiếu file mẫu, hãy copy vào data/sample/"
    result = read_and_validate_workbook(SAMPLE_PATH)
    assert result.is_valid, result.errors
    assert result.gop_vat_summary.total_data_rows > 0
    assert result.khop_ma_summary.total_data_rows > 0


def test_sample_file_keeps_leading_zero_codes():
    result = read_and_validate_workbook(SAMPLE_PATH)
    # Không có dòng nào bị mất số 0 đầu do bị Excel/openpyxl ép kiểu số.
    for m in result.mapping_rows:
        assert isinstance(m.ma_sp, str)


def test_missing_sheet_reports_error(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet khac"
    ws.append(["col1", "col2"])
    path = tmp_path / "thieu_sheet.xlsx"
    wb.save(path)

    result = read_and_validate_workbook(path)
    assert not result.is_valid
    assert any("Gộp VAT" in e for e in result.errors)
    assert any("Khớp mã" in e for e in result.errors)


def test_missing_required_column_reports_error(tmp_path):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Gộp VAT"
    ws1.append(["NGUỒN", "MÃ VAT KIOT"])  # thiếu nhiều cột bắt buộc
    ws2 = wb.create_sheet("Khớp mã")
    ws2.append(["Mã SP", "Tên SP", "Mã vat kiot", "Tên trước (Giá đầu  - Giá cuối) (VAT)"])
    path = tmp_path / "thieu_cot.xlsx"
    wb.save(path)

    result = read_and_validate_workbook(path)
    assert not result.is_valid
    assert any("thiếu cột bắt buộc" in e for e in result.errors)


def test_non_xlsx_file_rejected(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("a,b,c")
    result = read_and_validate_workbook(path)
    assert not result.is_valid
