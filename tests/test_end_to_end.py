from pathlib import Path

import openpyxl

from src.config import AppConfig
from src.excel_export import export_result_workbook, result_lines_to_tsv
from src.excel_reader import read_and_validate_workbook
from src.models import InvoiceInputRow
from src.pipeline import process_invoice


def _build_synthetic_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Gộp VAT"
    ws.append(
        [
            "NGUỒN", "Mã kiot đại diện", "Tên kiot đại diện", "MÃ VAT KIOT", "TÊN VAT KIOT",
            "TÊN KẾ TOÁN", "Đơn vị tính", "TỒN ĐẦU NGÀY", "GIÁ BÁN", "THUẾ VAT", "THUẾ NK",
            "TỔNG THUẾ MẤT", "TỔNG ĐANG NHÁP", "TỒN CÓ THỂ NHÁP",
        ]
    )
    ws.append(["MAY", "K1", "Kiot 1", "VAT10001", "Ấm giữ nhiệt VAT", "Ấm giữ nhiệt Hoa 1L",
               "Cái", 100000, 75000, 0.08, None, 0.08, 0, 120])
    ws.append(["CAO", "K2", "Kiot 2", "VAT10001", "Ấm giữ nhiệt VAT", "Ấm giữ nhiệt Hoa 1L",
               "Cái", 50000, 75000, 0.08, None, 0.08, 0, 80])
    ws.append(["MAY", "K3", "Kiot 3", "VAT10001", "Ấm giữ nhiệt VAT", "Ấm giữ nhiệt Hoa 1L (lô 2)",
               "Cái", 20000, 75000, 0.08, None, 0.08, 0, 0])  # tồn = 0, phải bị loại

    ws2 = wb.create_sheet("Khớp mã")
    ws2.append(["Mã SP", "Tên SP", "Mã vat kiot", "Tên trước (Giá đầu  - Giá cuối) (VAT)"])
    ws2.append(["APTH10", "Ấm giữ nhiệt Hoa 1L", "VAT10001", "Ấm giữ nhiệt 1000ML (72,100 - 78,500) (VAT)"])

    ws3 = wb.create_sheet("Yêu cầu")
    ws3.append(["Ghi chú tham khảo, không dùng làm dữ liệu"])

    wb.save(path)


def test_full_pipeline_end_to_end(tmp_path):
    workbook_path = tmp_path / "nguon.xlsx"
    _build_synthetic_workbook(workbook_path)

    check = read_and_validate_workbook(workbook_path)
    assert check.is_valid, check.errors

    config = AppConfig(source_tiebreak_priority=["MAY", "CAO"])

    invoice_rows = [
        InvoiceInputRow(input_row_id="r1", order_index=1, ma_sp="APTH10",
                         ten_sp="Ấm giữ nhiệt Hoa 1L", sl=150, gia_ban=75000),
    ]

    result = process_invoice(invoice_rows, check.stock_rows, check.mapping_rows, config)
    assert result.is_valid, result.errors
    assert sum(l.sl_phan_bo for l in result.result_lines) == 150

    export_path = tmp_path / "ket_qua.xlsx"
    export_result_workbook(
        export_path,
        result.result_lines,
        invoice_rows,
        source_file_name=check.file_name,
        source_file_hash=check.sha256,
        config_summary={"gia_nhap_decimal_places": config.gia_nhap_decimal_places},
    )
    assert export_path.exists()

    wb = openpyxl.load_workbook(export_path)
    assert set(["Kết quả", "Hóa đơn đầu vào", "Thông tin xử lý"]).issubset(set(wb.sheetnames))
    ws = wb["Kết quả"]
    assert ws.max_row >= 2  # header + ít nhất 1 dòng kết quả

    tsv = result_lines_to_tsv(result.result_lines)
    assert "\t" in tsv
    assert "<" not in tsv and ">" not in tsv  # không chứa HTML


def test_pipeline_blocks_when_stock_excludes_zero(tmp_path):
    workbook_path = tmp_path / "nguon2.xlsx"
    _build_synthetic_workbook(workbook_path)
    check = read_and_validate_workbook(workbook_path)

    config = AppConfig(source_tiebreak_priority=["MAY", "CAO"])
    # Yêu cầu SL vượt tổng tồn dương thực tế (120+80=200)
    invoice_rows = [
        InvoiceInputRow(input_row_id="r1", order_index=1, ma_sp="APTH10",
                         ten_sp="Ấm giữ nhiệt Hoa 1L", sl=500, gia_ban=75000),
    ]
    result = process_invoice(invoice_rows, check.stock_rows, check.mapping_rows, config)
    assert not result.is_valid
